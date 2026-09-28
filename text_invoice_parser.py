"""
Text-layer invoice parser.

Uses pdfplumber word positions to rebuild line-item rows from PDFs that
already have a text layer. Does not need Poppler or Tesseract.
"""

from __future__ import annotations

import logging
import re
from typing import List, Optional, Dict, Any

import pandas as pd
import pdfplumber

logger = logging.getLogger(__name__)

MONEY_RE = re.compile(
    r"^\$\s*-?\d{1,3}(?:,\d{3})*(?:\.\d{2})?$"
    r"|^\$\s*-?\d+(?:\.\d{2})?$"
    r"|^-?\d{1,3}(?:,\d{3})*\.\d{2}$"
    r"|^-?\d+\.\d{2}$"
)
QTY_RE = re.compile(r"^-?\d+(?:\.\d+)?$")
INV_NO_RE = re.compile(r"\b(INV[-\s]?\d{2,4}[-\s]?\d+)\b", re.I)
DATE_RE = re.compile(
    r"(Issue Date|Due Date|Invoice Date|Date)\s*:\s*([A-Za-z]+\s+\d{1,2},\s*\d{4}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})",
    re.I,
)
PO_RE = re.compile(r"\bP\.?O\.?\s*(?:Number|No\.?|#)\s*:?\s*([A-Z0-9-]+)", re.I)
TAX_RE = re.compile(r"(?:Estimated\s+)?Tax(?:\s*\(([^)]+)\))?\s*\$?\s*([\d,]+\.\d{2})", re.I)
SUBTOTAL_RE = re.compile(r"\bSubtotal\b\s*\$?\s*([\d,]+\.\d{2})", re.I)
TOTAL_RE = re.compile(r"\bTotal(?:\s+Due)?\b\s*\$?\s*([\d,]+\.\d{2})", re.I)

STOP_LABELS = (
    "subtotal",
    "total due",
    "amount due",
    "balance due",
    "estimated tax",
    "payment instructions",
    "thank you",
)
HEADER_HINTS = {"description", "qty", "quantity", "hours", "rate", "amount", "price", "unit"}


def _parse_money(token: str) -> Optional[float]:
    cleaned = token.replace("$", "").replace(",", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def _money_str(value: Optional[float]) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return f"{float(value):.2f}"


def _qty_value(value: Optional[float]) -> Optional[float | int]:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if float(value).is_integer():
        return int(value)
    return round(float(value), 2)


def _cluster_rows(words: List[Dict[str, Any]], y_tol: float = 3.5) -> List[List[Dict[str, Any]]]:
    if not words:
        return []
    ordered = sorted(words, key=lambda w: (w["top"], w["x0"]))
    rows: List[List[Dict[str, Any]]] = []
    current: List[Dict[str, Any]] = [ordered[0]]
    current_y = ordered[0]["top"]

    for word in ordered[1:]:
        if abs(word["top"] - current_y) <= y_tol:
            current.append(word)
        else:
            rows.append(sorted(current, key=lambda w: w["x0"]))
            current = [word]
            current_y = word["top"]
    rows.append(sorted(current, key=lambda w: w["x0"]))
    return rows


def _row_text(row: List[Dict[str, Any]]) -> str:
    return " ".join(w["text"] for w in row).strip()


def _is_stop_row(text: str) -> bool:
    lowered = text.lower().strip()
    return any(lowered.startswith(label) or f" {label}" in f" {lowered}" for label in STOP_LABELS)


def _is_header_row(text: str) -> bool:
    tokens = {t.strip(":/").lower() for t in re.split(r"\s+", text)}
    hits = sum(1 for hint in HEADER_HINTS if hint in tokens)
    return hits >= 2


def _split_row_cells(row: List[Dict[str, Any]], page_width: float) -> Dict[str, Any]:
    qty_x = page_width * 0.42
    rate_x = page_width * 0.58
    amount_x = page_width * 0.72

    desc_parts: List[str] = []
    qty = None
    rate = None
    amount = None

    for word in row:
        text = word["text"].strip()
        x0 = word["x0"]
        if x0 >= amount_x and MONEY_RE.match(text):
            amount = _parse_money(text)
        elif x0 >= rate_x and MONEY_RE.match(text):
            rate = _parse_money(text)
        elif x0 >= qty_x and QTY_RE.match(text) and "$" not in text:
            if qty is None:
                qty = float(text)
        else:
            desc_parts.append(text)

    return {
        "description": " ".join(desc_parts).strip(),
        "quantity": qty,
        "rate": rate,
        "amount": amount,
    }


def _looks_like_service_title(text: str) -> bool:
    if not text or text.endswith("."):
        return False
    words = text.split()
    if not (2 <= len(words) <= 10):
        return False
    return text[0].isupper() and not MONEY_RE.search(text)


def _extract_metadata(full_text: str) -> Dict[str, str]:
    meta = {
        "invoice_number": "",
        "vendor": "",
        "billed_to": "",
        "issue_date": "",
        "due_date": "",
        "po_number": "",
        "subtotal": "",
        "tax": "",
        "tax_rate": "",
        "total_due": "",
    }
    if not full_text:
        return meta

    lines = [ln.strip() for ln in full_text.splitlines() if ln.strip()]
    if lines:
        # First line is often "Vendor INVOICE"
        meta["vendor"] = re.sub(r"\bINVOICE\b", "", lines[0], flags=re.I).strip(" -|")

    inv = INV_NO_RE.search(full_text)
    if inv:
        meta["invoice_number"] = inv.group(1).replace(" ", "").upper()

    for label, value in DATE_RE.findall(full_text):
        key = "issue_date" if "issue" in label.lower() or label.lower() == "date" or "invoice" in label.lower() else "due_date"
        if "due" in label.lower():
            key = "due_date"
        meta[key] = value.strip()

    po = PO_RE.search(full_text)
    if po:
        meta["po_number"] = po.group(1).strip()

    # Billed-to block: text after BILLED TO until invoice details / address end.
    billed = re.search(
        r"BILLED\s+TO.*?\n(.*?)(?:HOURS\s*/|DESCRIPTION\b)",
        full_text,
        re.I | re.S,
    )
    if billed:
        block_lines = []
        for ln in billed.group(1).splitlines():
            cleaned = re.sub(
                r"(Issue Date|Due Date|Payment Terms|PO Number|Invoice Details).*$",
                "",
                ln,
                flags=re.I,
            ).strip(" ,")
            if cleaned and cleaned.upper() not in {"INVOICE DETAILS", "BILLED TO"}:
                block_lines.append(cleaned)
        meta["billed_to"] = ", ".join(block_lines[:4])

    sub = SUBTOTAL_RE.search(full_text)
    if sub:
        meta["subtotal"] = sub.group(1).replace(",", "")

    tax = TAX_RE.search(full_text)
    if tax:
        meta["tax_rate"] = (tax.group(1) or "").strip()
        meta["tax"] = tax.group(2).replace(",", "")

    total = TOTAL_RE.search(full_text)
    if total:
        meta["total_due"] = total.group(1).replace(",", "")

    return meta


class TextInvoiceParser:
    """Parse invoice line items from PDF text + word coordinates."""

    def __init__(self, debug: bool = False):
        self.debug = debug

    def extract_from_pdf(self, pdf_path: str, output_path: Optional[str] = None) -> pd.DataFrame:
        logger.info("Text-layer parse: %s", pdf_path)
        line_items: List[Dict[str, Any]] = []
        full_text_parts: List[str] = []

        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                full_text_parts.append(page.extract_text() or "")
                words = page.extract_words(use_text_flow=True, keep_blank_chars=False) or []
                if not words:
                    continue
                rows = _cluster_rows(words)
                line_items.extend(self._rows_to_items(rows, page.width))

        if not line_items:
            logger.warning("Text-layer parser found no line items")
            return pd.DataFrame()

        meta = _extract_metadata("\n".join(full_text_parts))
        formatted: List[Dict[str, Any]] = []
        for item in line_items:
            if item.get("amount") is None:
                continue
            formatted.append(
                {
                    "invoice_number": meta["invoice_number"],
                    "vendor": meta["vendor"],
                    "billed_to": meta["billed_to"],
                    "issue_date": meta["issue_date"],
                    "due_date": meta["due_date"],
                    "po_number": meta["po_number"],
                    "item": item.get("item", ""),
                    "details": item.get("details", ""),
                    "quantity": _qty_value(item.get("quantity")),
                    "rate": _money_str(item.get("rate")),
                    "amount": _money_str(item.get("amount")),
                    "subtotal": meta["subtotal"],
                    "tax": meta["tax"],
                    "tax_rate": meta["tax_rate"],
                    "total_due": meta["total_due"],
                }
            )

        df = pd.DataFrame(formatted)
        if df.empty:
            return df

        if output_path:
            self._save_output(df, output_path)

        return df

    def _save_output(self, df: pd.DataFrame, output_path: str) -> None:
        lower = output_path.lower()
        if lower.endswith(".xlsx"):
            self._save_xlsx(df, output_path)
        else:
            # utf-8-sig helps Excel show full Unicode text correctly.
            df.to_csv(output_path, index=False, encoding="utf-8-sig")
        logger.info("Saved table to %s", output_path)

    def _save_xlsx(self, df: pd.DataFrame, output_path: str) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
            from openpyxl.styles import Font, Alignment
        except ImportError:
            csv_fallback = output_path[:-5] + ".csv"
            df.to_csv(csv_fallback, index=False, encoding="utf-8-sig")
            logger.warning("openpyxl missing; wrote CSV instead: %s", csv_fallback)
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Invoice"

        headers = list(df.columns)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        for row in df.itertuples(index=False):
            ws.append(list(row))

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        widths = {
            "invoice_number": 16,
            "vendor": 24,
            "billed_to": 36,
            "issue_date": 18,
            "due_date": 18,
            "po_number": 14,
            "item": 40,
            "details": 55,
            "quantity": 10,
            "rate": 12,
            "amount": 12,
            "subtotal": 12,
            "tax": 12,
            "tax_rate": 10,
            "total_due": 12,
        }
        for idx, col in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, 14)

        wb.save(output_path)

    def _rows_to_items(self, rows: List[List[Dict[str, Any]]], page_width: float) -> List[Dict[str, Any]]:
        items: List[Dict[str, Any]] = []
        pending_title = ""
        in_table = False

        for row in rows:
            text = _row_text(row)
            if not text:
                continue

            if _is_header_row(text):
                in_table = True
                pending_title = ""
                continue

            if _is_stop_row(text):
                break

            if not in_table:
                continue

            cells = _split_row_cells(row, page_width)

            if cells["amount"] is not None:
                items.append(
                    {
                        "item": pending_title,
                        "details": cells["description"],
                        "quantity": cells["quantity"],
                        "rate": cells["rate"],
                        "amount": cells["amount"],
                    }
                )
                pending_title = ""
                continue

            desc = cells["description"] or text
            if _looks_like_service_title(desc):
                pending_title = desc
            elif items:
                prev = items[-1]["details"]
                items[-1]["details"] = f"{prev} {desc}".strip() if prev else desc
            elif desc:
                pending_title = f"{pending_title} {desc}".strip()

        return items
